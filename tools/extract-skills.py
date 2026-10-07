#!/usr/bin/env python3
"""
從 GM 的自動角色卡（xlsx）擷取技能目錄，輸出 src/data/skills.js。
只在更新技能資料時才需要跑；平常網站直接讀 src/data/skills.js。

用法（需要 openpyxl：pip install openpyxl）：
  python3 tools/extract-skills.py 標準卡.xlsx [補充卡1.xlsx 補充卡2.xlsx ...]
- 第一個檔案是標準：所有技能與每級數值都以它為準。
- 其餘檔案只用來「補」標準卡沒有的技能（個人專屬技能），標記 personal: true。
資料來源（每個 xlsx）：
  「技能表」A 名稱／B 位階／C 類型／D 系別／E 效果文字
  「後台_技能計算」AN:BL 的「技能 × 等級 → 累積數值加成」、AB:AL 的升級經驗表
輸出不含任何玩家資料。
"""
import json, re, sys
import openpyxl
from openpyxl.utils import column_index_from_string as CI

STAT_RENAME = {'體魄防禦': '體魄強韌', '抗性防禦': '抗性免疫'}  # 試算表與網站的名稱差異
TIERS = ['初階', '進階', '大師', '傳說']
ACTIVATE = re.compile(r'減少(\d+)點算力上限[^。]*。?[^。]*永久獲得啟動能力')  # 只有「永久獲得啟動能力」的一次性被動才算啟動類


def stat(name):
    return STAT_RENAME.get(name, name)


def read_catalog(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['技能表']
    info = {}
    for r in range(1, ws.max_row + 1):
        n = ws.cell(r, 1).value
        tier = ws.cell(r, 2).value
        if isinstance(n, str) and n.strip() and tier in TIERS + ['神級']:
            text = str(ws.cell(r, 5).value or '').replace('\r', '').strip()
            info[n.strip()] = {'tier': tier, 'kind': ws.cell(r, 3).value, 'school': ws.cell(r, 4).value, 'text': text}
    ws = wb['後台_技能計算']
    c0, c1 = CI('AN'), CI('BL')
    heads = [ws.cell(1, c).value for c in range(c0, c1 + 1)]
    fx = {}
    for r in range(2, ws.max_row + 1):
        n = ws.cell(r, c0).value
        lv = ws.cell(r, c0 + 1).value
        if n is None or lv is None:
            continue
        row = {}
        for h, c in zip(heads[2:], range(c0 + 2, c1 + 1)):
            v = ws.cell(r, c).value
            if h and isinstance(v, (int, float)) and v != 0:
                row[stat(h)] = int(v) if float(v).is_integer() else v
        fx.setdefault(str(n).strip(), {})[int(lv)] = row
    exp = {}
    for r in range(2, 8):
        t = ws.cell(r, CI('AB')).value
        if t in TIERS:
            exp[t] = [int(ws.cell(r, c).value) for c in range(CI('AC'), CI('AL') + 1)]
    return info, fx, exp


def main(paths):
    base_info, base_fx, exp = read_catalog(paths[0])
    table = {}
    for name, i in base_info.items():
        if name not in base_fx:
            print('警告：後台沒有數值表，略過', name)
            continue
        table[name] = dict(i, fx=[base_fx[name].get(l, {}) for l in range(1, 11)])
    for p in paths[1:]:
        info, fx, _ = read_catalog(p)
        for name, i in info.items():
            if name in table:
                continue
            if name in fx:
                table[name] = dict(i, fx=[fx[name].get(l, {}) for l in range(1, 11)], personal=True)
            else:  # 玩家自己加進技能表、後台沒有數值表：收進目錄但沒有自動數值（效果在該玩家的手動調整裡）
                table[name] = dict(i, fx=[{} for _ in range(10)], personal=True, manual=True)
            print('補入個人專屬技能', name, '（來自', p, '）', '' if name in fx else '＊後台沒有數值表，數值要手動調整')
    for name, t in table.items():
        m = ACTIVATE.search(t['text'])
        if m:  # 一次性被動：扣算力上限，之後才有「啟動」的數值（見 src/game/skillTable.js）
            t['activate'] = {'算力': int(m.group(1))}
    order = sorted(table, key=lambda n: (TIERS.index(table[n]['tier']) if table[n]['tier'] in TIERS else 9, n))
    out = ['// ============================================================',
           '// 技能目錄（自動產生，請勿手改；來源與更新方式見 tools/extract-skills.py）',
           '// fx[N-1] = 技能 N 級時「累積」的數值加成；啟動類技能（activate）要啟動後才算，且啟動時扣算力上限；manual = 沒有自動數值',
           '// ============================================================',
           f'export const EXP_TABLE = {json.dumps(exp, ensure_ascii=False)}; // 位階 → 升到 1~10 級各要多少經驗（試算表「技能等階」）',
           'export const SKILL_TABLE = {']
    for n in order:
        out.append(f'  {json.dumps(n, ensure_ascii=False)}: {json.dumps(table[n], ensure_ascii=False, separators=(",", ":"))},')
    out.append('};')
    open('src/data/skills.js', 'w', encoding='utf-8').write('\n'.join(out) + '\n')
    print('技能', len(table), '個；啟動類', sum('activate' in t for t in table.values()), '個；個人專屬', sum(bool(t.get('personal')) for t in table.values()), '個')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1:])
